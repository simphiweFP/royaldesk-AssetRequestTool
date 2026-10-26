#!/usr/bin/env python3
"""Import the existing E-Waste Asset Write-Off Register into RoyalDesk.

Usage:
  pip install openpyxl
  python tools/import-ewaste.py "E-Waste Asset Write-Off Register.xlsx" \
      --api https://localhost:7080 --username demo.user --password RoyalDesk123!

The workbook stays local. The script only posts mapped JSON rows to RoyalDesk.
"""
import argparse, base64, json, ssl, urllib.request
from datetime import date, datetime
from openpyxl import load_workbook


def as_iso(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return None if value in (None, "") else str(value)


def main():
    p = argparse.ArgumentParser()
    p.add_argument("workbook")
    p.add_argument("--api", default="https://localhost:7080")
    p.add_argument("--username", default="demo.user")
    p.add_argument("--password", default="RoyalDesk123!")
    p.add_argument("--verify-tls", action="store_true")
    args = p.parse_args()

    ws = load_workbook(args.workbook, data_only=True).active
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    required = ["Serial Number", "Make", "Model", "Category", "Condition", "Value", "Fully Depreciated", "Disposal Method", "Disposal Date", "Approved By", "Notes"]
    normalised = {h.strip(): i for i, h in enumerate(headers)}
    if "Value" not in normalised and "Value " in normalised:
        normalised["Value"] = normalised["Value "]
    missing = [h for h in required if h not in normalised]
    if missing:
        raise SystemExit(f"Missing columns: {', '.join(missing)}")

    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        serial = values[normalised["Serial Number"]]
        if serial in (None, ""):
            continue
        yes = str(values[normalised["Fully Depreciated"]] or "").strip().lower() in {"yes", "true", "1"}
        rows.append({
            "serialNumber": str(serial).strip(),
            "make": as_iso(values[normalised["Make"]]),
            "model": as_iso(values[normalised["Model"]]),
            "category": as_iso(values[normalised["Category"]]) or "Other",
            "condition": as_iso(values[normalised["Condition"]]) or "Obsolete",
            "value": values[normalised["Value"]],
            "fullyDepreciated": yes,
            "disposalMethod": as_iso(values[normalised["Disposal Method"]]),
            "disposalDate": as_iso(values[normalised["Disposal Date"]]),
            "approvedBy": as_iso(values[normalised["Approved By"]]),
            "notes": as_iso(values[normalised["Notes"]]),
        })

    token = base64.b64encode(f"{args.username}:{args.password}".encode()).decode()
    req = urllib.request.Request(
        args.api.rstrip("/") + "/api/lifecycle/imports/e-waste",
        data=json.dumps(rows).encode(),
        headers={"Content-Type": "application/json", "Authorization": f"Basic {token}"},
        method="POST",
    )
    context = None if args.verify_tls else ssl._create_unverified_context()
    with urllib.request.urlopen(req, context=context) as response:
        print(response.read().decode())


if __name__ == "__main__":
    main()
